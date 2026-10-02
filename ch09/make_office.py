"""09장 — 문서 파싱 도구 시험용 Office 문서(DOCX·PPTX·XLSX)를 평가셋 원본 내용으로 만듭니다.

내용을 코드로 넣으므로 정답(읽기 순서대로의 텍스트, 표)을 정확히 알 수 있습니다.
  office/report.docx  제목·문단·소제목 + 표 1개 (Word)
  office/slides.pptx  제목 슬라이드, 글머리 기호 슬라이드, 표 슬라이드 (PowerPoint)
  office/sheet.xlsx   표 2개를 시트 2개에 (Excel)
  office/office_gt.json  파일별 정답 텍스트와 표 (header + rows)
실행: python make_office.py   (pip install python-docx python-pptx openpyxl)
"""
import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
CONTENT = HERE.parent / "datasets" / "ko-ocr-bench" / "content"
OUT = HERE / "office"


def load(name):
    return json.loads((CONTENT / f"{name}.json").read_text(encoding="utf-8"))


def table_lines(table):
    """표를 정답 텍스트로: 머리글 칸, 이어서 행마다 칸을 차례로 한 줄씩."""
    return [str(c) for c in table["header"]] + [str(c) for row in table["rows"] for c in row]


def make_docx(doc, news, table):
    d = Document()
    lines = []
    d.add_heading(doc["title"], level=1)
    lines.append(doc["title"])
    for p in doc["paragraphs"]:
        d.add_paragraph(p)
        lines.append(p)
    d.add_heading(table["title"], level=2)
    lines.append(table["title"])
    t = d.add_table(rows=1 + len(table["rows"]), cols=len(table["header"]))
    t.style = "Table Grid"
    for r, row in enumerate([table["header"]] + table["rows"]):
        for c, value in enumerate(row):
            t.cell(r, c).text = str(value)
    lines += table_lines(table)
    article = news["articles"][0]
    d.add_heading(article["heading"], level=2)
    lines.append(article["heading"])
    for p in article["paragraphs"]:
        d.add_paragraph(p)
        lines.append(p)
    d.save(OUT / "report.docx")
    return lines


def make_pptx(doc, table):
    prs = Presentation()
    lines = []
    s = prs.slides.add_slide(prs.slide_layouts[0])  # 제목 슬라이드
    s.shapes.title.text = doc["title"]
    s.placeholders[1].text = "가상 안내 자료"
    lines += [doc["title"], "가상 안내 자료"]
    s = prs.slides.add_slide(prs.slide_layouts[1])  # 제목 + 글머리 기호
    s.shapes.title.text = "주요 내용"
    body = s.placeholders[1].text_frame
    body.text = doc["paragraphs"][0]
    for p in doc["paragraphs"][1:3]:
        body.add_paragraph().text = p
    lines += ["주요 내용"] + doc["paragraphs"][:3]
    s = prs.slides.add_slide(prs.slide_layouts[5])  # 제목만
    s.shapes.title.text = table["title"]
    rows = [table["header"]] + table["rows"]
    shape = s.shapes.add_table(len(rows), len(table["header"]), Inches(0.3), Inches(1.5), Inches(9.4), Inches(4))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            shape.table.cell(r, c).text = str(value)
    lines += [table["title"]] + table_lines(table)
    prs.save(OUT / "slides.pptx")
    return lines


def make_xlsx(tables):
    wb = Workbook()
    wb.remove(wb.active)
    lines = []
    for i, table in enumerate(tables, 1):
        ws = wb.create_sheet(f"표{i}")
        ws.append(table["header"])
        for row in table["rows"]:
            ws.append(row)
        lines += table_lines(table)
    wb.save(OUT / "sheet.xlsx")
    return lines


def main():
    OUT.mkdir(exist_ok=True)
    docs, tables, news = load("doc_ko"), load("table"), load("multicol")
    gt = {
        "report.docx": {"text": "\n".join(make_docx(docs[0], news[0], tables[0])), "tables": [tables[0]]},
        "slides.pptx": {"text": "\n".join(make_pptx(docs[1], tables[1])), "tables": [tables[1]]},
        "sheet.xlsx": {"text": "\n".join(make_xlsx(tables[2:4])), "tables": tables[2:4]},
    }
    (OUT / "office_gt.json").write_text(json.dumps(gt, ensure_ascii=False, indent=1), encoding="utf-8")
    for name, g in gt.items():
        print(f"{name:12s} 정답 {len(g['text']):5d}자, 표 {len(g['tables'])}개")


if __name__ == "__main__":
    main()
